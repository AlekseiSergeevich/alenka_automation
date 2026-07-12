import subprocess
import tempfile
import openpyxl
import os

with tempfile.TemporaryDirectory() as temp_dir:
    subprocess.run(["libreoffice", "--headless", "--nologo", "--nofirststartwizard", "--convert-to", "xlsx", "--outdir", temp_dir, "logs/order_store_170.xls"], capture_output=True)
    wb = openpyxl.load_workbook(os.path.join(temp_dir, "order_store_170.xlsx"))
    sh = wb.worksheets[0]
    # find header
    scan = min(80, sh.max_row)
    hdr_row = None
    for r in range(1, scan + 1):
        row_vals = {str(sh.cell(row=r, column=c).value).strip() for c in range(1, sh.max_column + 1) if sh.cell(row=r, column=c).value is not None}
        if {"УКП", "КОД Продаж"} <= row_vals:
            hdr_row = r
            break
            
    headers = [str(sh.cell(row=hdr_row, column=c).value) for c in range(1, sh.max_column + 1)]
    print("Headers:", headers)
    
    for r in range(1, sh.max_row + 1):
        row_vals = [str(sh.cell(row=r, column=c).value) for c in range(1, sh.max_column + 1)]
        if any("Аленка" in str(v) and "вареная" in str(v) for v in row_vals):
            print("Row:", r)
            for h, v in zip(headers, row_vals):
                if v != 'None':
                    print(f"{h}: {v}")
