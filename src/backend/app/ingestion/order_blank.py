"""Парсинг бланка заказа (.xls) и безопасный seed таблицы ``product``.

Seed только если ``product`` пустая; без удаления «лишних» артикулов (в отличие от CLI
``scripts/sync_products_from_order_blank.py`` в режиме полной синхронизации).
"""

from __future__ import annotations

import logging
from collections import Counter
from dataclasses import dataclass
from datetime import datetime, timezone
from enum import Enum
from pathlib import Path
from typing import Any

import xlrd
from sqlalchemy import func, select
from sqlalchemy.dialects.postgresql import insert as pg_insert
from sqlalchemy.ext.asyncio import AsyncSession

from src.backend.app.ingestion.utils import to_decimal
from src.backend.app.models import Product

logger = logging.getLogger(__name__)

DEFAULT_NAME_SUBSTR = "Бланк заказа"
_EXPECTED_MARKERS = frozenset({"УКП", "Название SKU", "КОД Продаж"})


class OrderBlankSeedStatus(str, Enum):
    skipped_nonempty = "skipped_nonempty"
    skipped_no_file = "skipped_no_file"
    skipped_parse_empty = "skipped_parse_empty"
    inserted = "inserted"


@dataclass(slots=True)
class OrderBlankSeedResult:
    status: OrderBlankSeedStatus
    rows_inserted: int = 0
    path: Path | None = None


def project_root_from_here() -> Path:
    """Корень репозитория: .../src/backend/app/ingestion/order_blank.py -> parents[4]."""

    return Path(__file__).resolve().parents[4]


def find_latest_order_blank_xls(project_root: Path, name_substr: str = DEFAULT_NAME_SUBSTR) -> Path | None:
    raw_dir = project_root / "data" / "raw"
    if not raw_dir.is_dir():
        return None
    candidates = []
    for ext in ("*.xls", "*.xlsx"):
        candidates.extend(p for p in raw_dir.glob(ext) if name_substr in p.name)
    candidates.sort(key=lambda p: p.stat().st_mtime, reverse=True)
    return candidates[0] if candidates else None


def load_2d_grid(path: Path) -> list[list[str]]:
    if path.suffix.lower() == ".xlsx":
        import openpyxl
        wb = openpyxl.load_workbook(str(path), data_only=True)
        sh = None
        for name in wb.sheetnames:
            if "Бланк заказа" in name:
                sh = wb[name]
                break
        if sh is None:
            sh = wb.active
        grid = []
        for row in sh.iter_rows(values_only=True):
            grid.append([str(v).strip() if v is not None else "" for v in row])
        return grid
    else:
        import xlrd
        wb = xlrd.open_workbook(str(path))
        try:
            sh = wb.sheet_by_name("Бланк заказа")
        except xlrd.XLRDError:
            sh = wb.sheet_by_index(0)
        grid = []
        for r in range(sh.nrows):
            row_data = []
            for c in range(sh.ncols):
                v = sh.cell_value(r, c)
                typ = sh.cell_type(r, c)
                if typ == xlrd.XL_CELL_EMPTY:
                    row_data.append("")
                elif typ == xlrd.XL_CELL_NUMBER:
                    if isinstance(v, float) and v == int(v):
                        row_data.append(str(int(v)))
                    else:
                        row_data.append(str(v).strip())
                else:
                    row_data.append(str(v).strip())
            grid.append(row_data)
        return grid


def _cell_str(grid: list[list[str]], r: int, c: int) -> str:
    if r < 0 or r >= len(grid):
        return ""
    if c < 0 or c >= len(grid[r]):
        return ""
    return grid[r][c]


def _find_header_row(grid: list[list[str]]) -> int:
    scan = min(80, len(grid))
    for r in range(scan):
        cells = set(grid[r])
        if _EXPECTED_MARKERS <= cells:
            return r
    msg = "Не найдена строка заголовка с колонками УКП / Название SKU / КОД Продаж"
    raise ValueError(msg)


def _build_header_map(grid: list[list[str]], hdr_row: int) -> dict[str, int]:
    out: dict[str, int] = {}
    row = grid[hdr_row]
    for c, val in enumerate(row):
        if val:
            out[val] = c
    need = (
        "УКП",
        "КОД Продаж",
        "Название SKU",
        "Тип",
        "Фокус",
        "Группа ABC",
        "Признаки",
        "В кор. шт/кг",
        "Бренд",
        "Базовая цена за короб",
        "Срок годности, дн.",
        "Фасовка - вес штуки, гр",
    )
    missing = [k for k in need if k not in out]
    if missing:
        msg = f"В заголовке нет колонок: {', '.join(missing)}"
        raise ValueError(msg)
    return out


def _article_for_row(kod: str, ukp: str, ukp_counts: Counter[str]) -> str:
    kod_t = kod.strip()
    ukp_t = ukp.strip()
    if ukp_counts[ukp_t] > 1:
        return kod_t or ukp_t
    return ukp_t


def parse_order_blank(path: Path) -> tuple[list[dict[str, Any]], list[str]]:
    grid = load_2d_grid(path)

    hdr = _find_header_row(grid)
    col = _build_header_map(grid, hdr)

    ukp_i = col["УКП"]
    kod_i = col["КОД Продаж"]

    ukp_counts: Counter[str] = Counter()
    for r in range(hdr + 1, len(grid)):
        ukp = _cell_str(grid, r, ukp_i)
        if ukp:
            ukp_counts[ukp] += 1

    now = datetime.now(timezone.utc)
    rows_out: list[dict[str, Any]] = []

    for r in range(hdr + 1, len(grid)):
        ukp = _cell_str(grid, r, ukp_i)
        if not ukp:
            continue
        kod = _cell_str(grid, r, kod_i)
        article = _article_for_row(kod, ukp, ukp_counts).strip()
        if not article:
            continue

        qty_raw = _cell_str(grid, r, col["В кор. шт/кг"])
        qty = to_decimal(qty_raw)

        rows_out.append(
            {
                "article": article[:128],
                "nom_number": None,
                "name": _cell_str(grid, r, col["Название SKU"])[:1024] or article[:1024],
                "unit": "",
                "type": _cell_str(grid, r, col["Тип"])[:64],
                "focus": _cell_str(grid, r, col["Фокус"])[:64],
                "group_abc": _cell_str(grid, r, col["Группа ABC"])[:64],
                "feature": _cell_str(grid, r, col["Признаки"])[:64],
                "quantity_in_box": qty,
                "brand": _cell_str(grid, r, col["Бренд"])[:128],
                "base_price": to_decimal(_cell_str(grid, r, col["Базовая цена за короб"])),
                "shelf_life_days": to_decimal(_cell_str(grid, r, col["Срок годности, дн."])),
                "weight_gr": to_decimal(_cell_str(grid, r, col["Фасовка - вес штуки, гр"])),
                "updated_at": now,
            }
        )

    dedup: dict[str, dict[str, Any]] = {}
    for row in rows_out:
        dedup[row["article"]] = row

    merged = list(dedup.values())
    arts = list(dedup.keys())

    return merged, arts


async def product_row_count(session: AsyncSession) -> int:
    n = await session.scalar(select(func.count()).select_from(Product))
    return int(n or 0)


async def seed_products_from_order_blank_if_empty(
    session: AsyncSession,
    path: Path | None,
    *,
    project_root: Path | None = None,
) -> OrderBlankSeedResult:
    """Вставить товары из бланка только если таблица ``product`` пустая.

    Конфликты по ``article`` игнорируются (не перезаписываем существующие строки).
    """

    if await product_row_count(session) > 0:
        return OrderBlankSeedResult(status=OrderBlankSeedStatus.skipped_nonempty)

    root = project_root or project_root_from_here()
    resolved = path
    if resolved is None:
        resolved = find_latest_order_blank_xls(root)
    if resolved is None or not resolved.is_file():
        logger.info("Order blank seed skipped: no file at %s", path or "(auto)")
        return OrderBlankSeedResult(status=OrderBlankSeedStatus.skipped_no_file)

    rows, articles = parse_order_blank(resolved.resolve())
    if not articles:
        return OrderBlankSeedResult(
            status=OrderBlankSeedStatus.skipped_parse_empty,
            path=resolved.resolve(),
        )

    now = datetime.now(timezone.utc)
    product_keys = {"article", "name", "unit", "type", "focus", "group_abc", "feature", "quantity_in_box", "nom_number", "updated_at"}
    product_rows = []
    for row in rows:
        p_row = {k: v for k, v in row.items() if k in product_keys}
        p_row["updated_at"] = now
        product_rows.append(p_row)

    stmt = pg_insert(Product).values(product_rows)
    stmt = stmt.on_conflict_do_nothing(index_elements=[Product.article])
    await session.execute(stmt)

    return OrderBlankSeedResult(
        status=OrderBlankSeedStatus.inserted,
        rows_inserted=len(rows),
        path=resolved.resolve(),
    )
