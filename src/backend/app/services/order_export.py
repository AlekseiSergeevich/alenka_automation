import os
import subprocess
import tempfile
import math
import shutil
from pathlib import Path
from collections import Counter
import openpyxl

def export_order_to_xls(template_path: str | Path, output_path: str | Path, order_data: dict[str, float]) -> None:
    """
    order_data: dict mapping article to quantity (in kg)
    """
    template_path = str(Path(template_path).resolve())
    output_path = str(Path(output_path).resolve())
    
    with tempfile.TemporaryDirectory() as temp_dir:
        # 1. Конвертируем шаблон .xls в .xlsx через LibreOffice, чтобы сохранить формулы
        res = subprocess.run(
            ["libreoffice", "--headless", "--nologo", "--nofirststartwizard", "--convert-to", "xlsx", "--outdir", temp_dir, template_path],
            capture_output=True, text=True
        )
        if res.returncode != 0:
            raise RuntimeError(f"Failed to convert template to xlsx: {res.stderr}\n{res.stdout}")
            
        temp_xlsx = os.path.join(temp_dir, Path(template_path).stem + ".xlsx")
        if not os.path.exists(temp_xlsx):
            raise RuntimeError("LibreOffice returned success but xlsx file not found")
        
        # 2. Модифицируем файл через openpyxl
        wb = openpyxl.load_workbook(temp_xlsx)
        
        # Locate the target sheet
        if "Бланк заказа" in wb.sheetnames:
            sh = wb["Бланк заказа"]
        else:
            sh = wb.worksheets[0]
            
        # Find header row (1-indexed in openpyxl)
        scan = min(80, sh.max_row)
        hdr_row = None
        for r in range(1, scan + 1):
            row_vals = {str(sh.cell(row=r, column=c).value).strip() for c in range(1, sh.max_column + 1) if sh.cell(row=r, column=c).value is not None}
            if {"УКП", "КОД Продаж"} <= row_vals:
                hdr_row = r
                break
                
        if hdr_row is None:
            raise ValueError("Header row not found")
            
        col_map = {}
        for c in range(1, sh.max_column + 1):
            val = sh.cell(row=hdr_row, column=c).value
            if val is not None:
                key = str(val).strip()
                if key:
                    col_map[key] = c
                    
        target_col = col_map.get("Заявка, короба") or col_map.get("Заявка, шт.")
        if target_col is None:
            raise ValueError("Could not find order quantity column (Заявка, короба / Заявка, шт.)")
            
        weight_col = col_map.get("Вес короба, кг")
        if weight_col is None:
            raise ValueError("Could not find column 'Вес короба, кг'")
            
        ukp_i = col_map["УКП"]
        kod_i = col_map.get("КОД Продаж", ukp_i)
        
        # Calculate duplicates to map rows properly
        ukp_counts = Counter()
        for r in range(hdr_row + 1, sh.max_row + 1):
            val = sh.cell(row=r, column=ukp_i).value
            if val is not None:
                ukp = str(val).strip()
                if ukp:
                    ukp_counts[ukp] += 1
                    
        # Iterate and write
        for r in range(hdr_row + 1, sh.max_row + 1):
            val_ukp = sh.cell(row=r, column=ukp_i).value
            if val_ukp is None:
                continue
                
            ukp = str(val_ukp).strip()
            if not ukp:
                continue
                
            val_kod = sh.cell(row=r, column=kod_i).value
            kod = str(val_kod).strip() if val_kod is not None else ""
            
            # determine article
            if ukp_counts[ukp] > 1:
                article = kod or ukp
            else:
                article = ukp
                
            if article in order_data:
                qty_kg = float(order_data[article])
                if qty_kg > 0:
                    weight_val = sh.cell(row=r, column=weight_col).value
                    # Parse weight
                    box_weight = 0.0
                    try:
                        box_weight = float(str(weight_val).replace(",", ".").strip())
                    except (ValueError, TypeError):
                        pass
                        
                    if box_weight > 0:
                        qty_boxes = math.ceil(qty_kg / box_weight)
                    else:
                        qty_boxes = math.ceil(qty_kg)
                        
                    sh.cell(row=r, column=target_col, value=qty_boxes)
                    
        modified_xlsx = os.path.join(temp_dir, "modified.xlsx")
        wb.save(modified_xlsx)
        
        # 3. Конвертируем обратно в .xls
        res = subprocess.run(
            ["libreoffice", "--headless", "--nologo", "--nofirststartwizard", "--convert-to", "xls", "--outdir", temp_dir, modified_xlsx],
            capture_output=True, text=True
        )
        if res.returncode != 0:
            raise RuntimeError(f"Failed to convert modified file back to xls: {res.stderr}\n{res.stdout}")
            
        final_xls = os.path.join(temp_dir, "modified.xls")
        if not os.path.exists(final_xls):
            raise RuntimeError("LibreOffice returned success but final xls file not found")
            
        # Move to output path
        shutil.move(final_xls, output_path)
